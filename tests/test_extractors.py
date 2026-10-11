import io
import zipfile
import pytest
from app.processing.extractors.text_extractor import TextExtractor
from app.processing.extractors.html_extractor import HTMLExtractor
from app.processing.extractors.eml_extractor import EMLExtractor
from app.processing.extractors.zip_handler import SafeZipHandler
from app.core.exceptions import FileValidationError


def test_text_extractor():
    extractor = TextExtractor()
    sample = "Hello World! This is plain text document."
    res = extractor.extract(sample.encode("utf-8"), "test.txt")
    assert "Hello World!" in res.text
    assert len(res.pages) == 1


def test_html_extractor():
    extractor = HTMLExtractor()
    html = "<html><head><title>Test Page</title></head><body><h1>Heading</h1><p>Paragraph text.</p></body></html>"
    res = extractor.extract(html.encode("utf-8"), "test.html")
    assert "Heading" in res.text
    assert "Paragraph text." in res.text


def test_eml_extractor():
    extractor = EMLExtractor()
    eml_raw = b"From: sender@acme.com\nTo: recv@acme.com\nSubject: Test Email\n\nEmail body content goes here."
    res = extractor.extract(eml_raw, "test.eml")
    assert "Test Email" in res.text
    assert "Email body content" in res.text


def test_safe_zip_extraction():
    zip_handler = SafeZipHandler()

    # Create in-memory valid zip
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("doc1.txt", "Document 1 content")
        zf.writestr("doc2.txt", "Document 2 content")
    zip_bytes = buf.getvalue()

    files = zip_handler.extract_files(zip_bytes)
    assert len(files) == 2
    filenames = [f[0] for f in files]
    assert "doc1.txt" in filenames
    assert "doc2.txt" in filenames


def test_zip_slip_prevention():
    zip_handler = SafeZipHandler()

    # Create zip with malicious Zip Slip path traversal member "../../outside.txt"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("../../outside.txt", "Malicious content")
    zip_bytes = buf.getvalue()

    with pytest.raises(FileValidationError):
        zip_handler.extract_files(zip_bytes)


def test_docx_extractor():
    import docx
    from app.processing.extractors.docx_extractor import DOCXExtractor

    doc = docx.Document()
    doc.add_paragraph("Enterprise Strategy Q4 Report.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Revenue"
    table.rows[0].cells[1].text = "$100M"
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    extractor = DOCXExtractor()
    res = extractor.extract(docx_bytes, "report.docx")
    assert "Enterprise Strategy Q4 Report." in res.text
    assert "Revenue | $100M" in res.text
    assert res.metadata["paragraph_count"] >= 1
    assert res.metadata["table_count"] == 1


def test_pptx_extractor():
    from pptx import Presentation
    from app.processing.extractors.pptx_extractor import PPTXExtractor

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    title.text = "Q4 Executive Presentation"
    buf = io.BytesIO()
    prs.save(buf)
    pptx_bytes = buf.getvalue()

    extractor = PPTXExtractor()
    res = extractor.extract(pptx_bytes, "deck.pptx")
    assert "Q4 Executive Presentation" in res.text
    assert res.metadata["slide_count"] == 1


def test_xlsx_extractor():
    import openpyxl
    from app.processing.extractors.xlsx_extractor import XLSXExtractor

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Financials"
    ws.append(["Metric", "Value"])
    ws.append(["Gross Margin", "85%"])
    buf = io.BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()

    extractor = XLSXExtractor()
    res = extractor.extract(xlsx_bytes, "sheets.xlsx")
    assert "Gross Margin | 85%" in res.text
    assert res.metadata["sheet_count"] == 1


def test_pdf_extractor():
    import fitz
    from app.processing.extractors.pdf_extractor import PDFExtractor

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Confidential Enterprise Audit Blueprint and Security Baseline.")
    pdf_bytes = doc.tobytes()

    extractor = PDFExtractor()
    res = extractor.extract(pdf_bytes, "audit.pdf")
    assert "Confidential Enterprise Audit Blueprint" in res.text
    assert res.metadata["page_count"] == 1


def test_extraction_service_routing():
    from app.services.extraction_service import ExtractionService

    service = ExtractionService()
    res = service.process_file(b"This is a plain markdown document with English words.", "notes.md")
    assert "plain markdown document" in res.text
    assert res.metadata["detected_language"] in ("en", "unknown")

