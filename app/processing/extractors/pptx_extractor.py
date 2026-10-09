import io
from pptx import Presentation
from app.processing.extractors.base import BaseExtractor, ExtractionResult, PageContent


class PPTXExtractor(BaseExtractor):
    def extract(self, file_bytes: bytes, filename: str) -> ExtractionResult:
        prs = Presentation(io.BytesIO(file_bytes))
        pages: list[PageContent] = []
        full_text_list: list[str] = []

        for slide_idx, slide in enumerate(prs.slides):
            slide_num = slide_idx + 1
            slide_texts: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_texts.append(shape.text.strip())

            combined_slide = "\n".join(slide_texts)
            pages.append(PageContent(page_number=slide_num, text=combined_slide))
            if combined_slide:
                full_text_list.append(f"--- Slide {slide_num} ---\n{combined_slide}")

        return ExtractionResult(
            text="\n\n".join(full_text_list),
            pages=pages,
            metadata={"slide_count": len(prs.slides)},
            extraction_method="pptx_python_pptx",
        )
