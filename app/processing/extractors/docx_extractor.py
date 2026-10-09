import io
import docx
from app.processing.extractors.base import BaseExtractor, ExtractionResult, PageContent


class DOCXExtractor(BaseExtractor):
    def extract(self, file_bytes: bytes, filename: str) -> ExtractionResult:
        doc = docx.Document(io.BytesIO(file_bytes))
        text_parts: list[str] = []

        # Extract paragraphs
        for p in doc.paragraphs:
            if p.text.strip():
                text_parts.append(p.text.strip())

        # Extract table content
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text_parts.append(" | ".join(row_text))

        combined_text = "\n\n".join(text_parts)
        page = PageContent(page_number=1, text=combined_text, ocr_used=False)

        return ExtractionResult(
            text=combined_text,
            pages=[page],
            metadata={"paragraph_count": len(doc.paragraphs), "table_count": len(doc.tables)},
            extraction_method="docx_python_docx",
        )
