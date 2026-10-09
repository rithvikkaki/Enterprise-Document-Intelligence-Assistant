import io
import openpyxl
from app.processing.extractors.base import BaseExtractor, ExtractionResult, PageContent


class XLSXExtractor(BaseExtractor):
    def extract(self, file_bytes: bytes, filename: str) -> ExtractionResult:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        pages: list[PageContent] = []
        full_text_list: list[str] = []

        for sheet_idx, sheet_name in enumerate(wb.sheetnames):
            sheet = wb[sheet_name]
            sheet_rows: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                row_vals = [str(val).strip() for val in row if val is not None and str(val).strip()]
                if row_vals:
                    sheet_rows.append(" | ".join(row_vals))

            combined_sheet = "\n".join(sheet_rows)
            pages.append(PageContent(page_number=sheet_idx + 1, text=combined_sheet))
            if combined_sheet:
                full_text_list.append(f"--- Sheet: {sheet_name} ---\n{combined_sheet}")

        return ExtractionResult(
            text="\n\n".join(full_text_list),
            pages=pages,
            metadata={"sheet_count": len(wb.sheetnames)},
            extraction_method="xlsx_openpyxl",
        )
